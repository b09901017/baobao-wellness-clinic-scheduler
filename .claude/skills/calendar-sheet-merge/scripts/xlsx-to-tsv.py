#!/usr/bin/env python3
"""舊試算表 .xlsx → 一張分頁一個 TSV。**不去識別化**，所以輸出只能放暫存區。

repo 裡另有一支 `scripts/legacy-xlsx-to-tsv.py`，那一支會把姓名與手寫註記抹掉，
用途是產生可以進版控的樣本。這一支相反：比對要拿真名去對行事曆，一個字都不能改。

所以它有一條硬規則：**輸出資料夾必須在 repo 外面**（預設是本次工作階段的暫存區）。
真實客戶資料進了 git 歷史就拿不掉（SPEC 第 10 節）。

用法：
    python3 xlsx-to-tsv.py <原檔.xlsx> <輸出資料夾>
"""

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('需要 openpyxl：pip install openpyxl')

REPO = Path(__file__).resolve().parents[4]


def cell_text(v):
    """讀成她從畫面上複製貼上的樣子 —— 匯入器就是照那個寫的。"""
    if v is None:
        return ''
    if v is True or v is False:
        return 'TRUE' if v else 'FALSE'
    if hasattr(v, 'strftime'):
        # 表頭的日期在畫面上顯示成「6月15日」，沒有年份。這裡不要自作主張補上，
        # 補哪一年是比對那一步要決定並且講出來的事。
        return f'{v.month}月{v.day}日'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src, out = Path(sys.argv[1]), Path(sys.argv[2]).resolve()

    if REPO in out.parents or out == REPO:
        sys.exit(f'拒絕寫進 repo（{out}）。真實客戶資料不進版控，請輸出到暫存區。')

    out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src, data_only=True)
    written = []
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':
            continue
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
            rows.append('\t'.join(
                cell_text(c.value).replace('\n', '\\n').replace('\t', ' ') for c in row
            ))
        while rows and not rows[-1].strip():
            rows.pop()
        (out / f'{ws.title}.tsv').write_text('\n'.join(rows), encoding='utf-8')
        written.append(ws.title)

    print(f'{len(written)} 張分頁 → {out}')
    for name in written:
        print(f'  {name}')


if __name__ == '__main__':
    main()
