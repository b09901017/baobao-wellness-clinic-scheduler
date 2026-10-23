# 讀法.json → 一人一頁的月曆試算表。
#
#   python build.py <月份資料夾>        例：.local/references/客戶時間/2026-11
#
# 讀 <月份資料夾>/過程/讀法.json（與 客戶脈絡.json，拿體內金屬），寫 <月份資料夾>/客戶時間-YYYY-MM.xlsx。
# 版面是她 2026-10-23 一版一版改出來的：一人一頁、一個月曆、紙色底、名字底下一塊「客戶交代（原文）」、
# 體內金屬一個紅色的標。規則（哪一天是什麼顏色）只從讀法.json 來；這支只負責畫。核對在 verify.py。
import sys, json, datetime, calendar, os
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WEEKDAY = {'一': 0, '二': 1, '三': 2, '四': 3, '五': 4, '六': 5, '日': 6}

PAPER = 'F3F0EA'  # 整頁的底色（她 10/23：太白、太淺淡會看不清楚）
# 種類 → (底色, 字色, 格子裡的字, 圖例)。verify.py 照底色認回來，改顏色要兩邊一起改
LOOK = {
    'no':           ('F4B9B3', '962A22', '不行',     '不行'),
    'am':           ('F6D08F', '6E4304', '上午不行', '只有部分時間可以'),
    'pm':           ('F6D08F', '6E4304', '下午不行', '只有部分時間可以'),
    'partial':      ('F6D08F', '6E4304', '',         '只有部分時間可以'),
    'yes':          ('B9E2C6', '1D6439', '可以',     '可以，優先排'),
    'free':         ('B9E2C6', '1D6439', '可以',     '可以，優先排'),
    'unsure':       ('F9E07E', '665000', '不一定',   '不一定'),
    'notMentioned': ('E4E0D8', '9A948A', '',         '沒提到'),
}
LEGEND_ORDER = ['可以，優先排', '不行', '只有部分時間可以', '不一定', '沒提到']
# 後面的蓋前面的：一天同時被說「可以」又「不行」時，不行贏
PRIORITY = ['yes', 'partial', 'am', 'pm', 'unsure', 'no']
# 客戶身上的警示裡，要在頁面上標紅的（她 10/23：體內金屬要標紅）
RED_FLAGS = ['金屬']


def expand(entries, year, month):
    """[3, '9-13', '週一', '週末'] → 那個月的日子。認不得的丟錯（不猜）。"""
    last = calendar.monthrange(year, month)[1]
    out = []
    for e in entries:
        if isinstance(e, int):
            days = [e]
        elif e == '週末':
            days = [d for d in range(1, last + 1) if datetime.date(year, month, d).weekday() >= 5]
        elif isinstance(e, str) and e.startswith('週') and len(e) == 2 and e[1] in WEEKDAY:
            days = [d for d in range(1, last + 1) if datetime.date(year, month, d).weekday() == WEEKDAY[e[1]]]
        elif e == '整個月':
            days = list(range(1, last + 1))
        elif isinstance(e, str) and '-' in e:
            a, b = (int(x) for x in e.split('-'))
            days = list(range(a, b + 1))
        else:
            raise ValueError(f'認不得的日子寫法：{e!r}（只收 數字、"9-13"、"週一"…"週日"、"週末"、"整個月"）')
        for d in days:
            if not 1 <= d <= last:
                raise ValueError(f'{e!r} 超出這個月（1～{last}）')
        out += days
    return out


def cells_of(person, year, month):
    last = calendar.monthrange(year, month)[1]
    rest = person.get('rest', 'free')
    if rest not in ('free', 'notMentioned'):
        raise ValueError(f"{person['name']}：rest 只能是 free 或 notMentioned")
    cells = {d: rest for d in range(1, last + 1)}
    days = person.get('days', {})
    for k in days:
        if k not in PRIORITY:
            raise ValueError(f"{person['name']}：days 裡認不得的種類 {k!r}")
    for kind in PRIORITY:
        for d in expand(days.get(kind, []), year, month):
            cells[d] = kind
    return cells


def word_of(person, kind, day, year, month):
    words = person.get('dayWords', {})
    for key, text in words.items():
        if key.isdigit() and int(key) == day:
            return text
    # 「週二」「10-30」這種一次寫好幾天的字，只寫到有標顏色、又不是「不行」的那幾天：
    # 「週二 3:30 後才行」不可以印在 9～13 整段不行裡的那個週二上；「9-12 選一天」也不可以寫到中間沒提到的 11 號
    if kind not in ('no', 'free', 'notMentioned'):
        for key, text in words.items():
            if not key.isdigit() and day in expand([key], year, month):
                return text
    return LOOK[kind][2]


def red_flags_by_name(folder):
    """客戶脈絡.json → {名字: [要標紅的警示]}。對不到剛好一位客戶的不標（也不猜），回報給呼叫端。"""
    path = os.path.join(folder, '過程', '客戶脈絡.json')
    if not os.path.exists(path):
        return None
    out = {}
    for p in json.load(open(path, encoding='utf-8'))['people']:
        if len(p['customers']) != 1:
            continue
        flags = p['customers'][0].get('flags') or []
        out[p['name']] = [f for f in flags if any(k in str(f) for k in RED_FLAGS)]
    return out


def chart_by_name(folder):
    """客戶脈絡.json → {名字: (病歷號那一格要寫的字, 是不是要提醒她)}。

    她 10/23：「那張表也要呈現所有人的病歷號，很重要」。畫的是她存的樣子（`chartNo`，前面的 0 照留）；
    那一格空著、號碼還寫在短備註裡的，用 app 認人的那一支（`chartNosOf()`）讀出來的。
    沒有、或同一位有兩個不一樣的號碼 → 照實寫、標成要提醒（不挑一個）。對不到剛好一位客戶的不寫。
    """
    path = os.path.join(folder, '過程', '客戶脈絡.json')
    if not os.path.exists(path):
        return None
    out = {}
    for p in json.load(open(path, encoding='utf-8'))['people']:
        if len(p['customers']) != 1:
            continue
        c = p['customers'][0]
        nos = c.get('chartNos') or []
        raw = str(c.get('chartNo') or '').strip()
        if len(nos) > 1:
            out[p['name']] = ('／'.join(nos), True)
        elif raw:
            out[p['name']] = (raw, False)
        elif nos:
            out[p['name']] = (nos[0], False)
        else:
            out[p['name']] = ('app 上沒有', True)
    return out


def lines_needed(text, size, width_units):
    """一行放不下時要幾行（中文字大約 size/5.5 個欄寬單位）。只用來估列高。"""
    per_char = size / 5.5
    return max(1, -(-int(len(text) * per_char) // int(width_units)))


def build(folder):
    plan = json.load(open(os.path.join(folder, '過程', '讀法.json'), encoding='utf-8'))
    year, month = (int(x) for x in plan['month'].split('-'))
    last = calendar.monthrange(year, month)[1]
    red = red_flags_by_name(folder)
    charts = chart_by_name(folder)
    if red is None:
        print('注意：沒有 過程/客戶脈絡.json —— 體內金屬與病歷號都放不上去，先跑 history.mjs')

    FONT = 'Microsoft JhengHei'
    INK, SOFT = '2B3440', '4B5563'
    def f(size, bold=False, color=INK): return Font(name=FONT, size=size, bold=bold, color=color)
    def fill(c): return PatternFill('solid', start_color=c, end_color=c)
    GAP = Side(style='thick', color=PAPER)
    COLS = list(range(2, 9))  # B–H
    BOX_WIDTH = 13 * 7

    first = datetime.date(year, month, 1)
    first_monday = first - datetime.timedelta(days=first.weekday())
    end = datetime.date(year, month, last)
    weeks = ((end - first_monday).days // 7) + 1

    wb = Workbook()
    wb.remove(wb.active)
    names = set()
    unflagged = []
    for p in plan['people']:
        title = p['name'][:31]
        if title in names:
            raise ValueError(f'兩頁同名：{title}')
        names.add(title)
        cells = cells_of(p, year, month)
        for key in p.get('dayWords', {}):
            if key.isdigit() and cells[int(key)] in ('free', 'notMentioned'):
                raise ValueError(f"{p['name']}：{key} 號寫了字，但那一天沒有標任何種類")
            if key.isdigit() and cells[int(key)] == 'no':
                raise ValueError(f"{p['name']}：{key} 號是「不行」，格子裡只能寫不行 —— 把那一個字拿掉，或改那一天的種類")

        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 110
        ws.column_dimensions['A'].width = 4
        for c in COLS:
            ws.column_dimensions[get_column_letter(c)].width = 13
        ws.column_dimensions['I'].width = 4

        ws.row_dimensions[1].height = 26
        ws.merge_cells('G1:H1')
        ws['G1'] = f'{year} 年 {month} 月'
        ws['G1'].font = f(11, False, SOFT)
        ws['G1'].alignment = Alignment(horizontal='right', vertical='bottom')
        ws.row_dimensions[2].height = 40
        ws.merge_cells('B2:D2')
        ws['B2'] = p['name']
        ws['B2'].font = f(24, True)
        ws['B2'].alignment = Alignment(vertical='bottom')
        # 病歷號：名字右邊，數字大、上面一行小字標「病歷號」
        chart = (charts or {}).get(p['name'])
        if chart:
            text, warn = chart
            ws.merge_cells('E1:F1')
            ws['E1'] = '病歷號'
            ws['E1'].font = f(10, True, SOFT)
            ws['E1'].alignment = Alignment(vertical='bottom', indent=1)
            ws.merge_cells('E2:F2')
            ws['E2'] = text
            ws['E2'].font = f(14 if warn else 20, True, 'B45309' if warn else INK)
            ws['E2'].alignment = Alignment(vertical='bottom', indent=1)
        flags = (red or {}).get(p['name'])
        if red is not None and p['name'] not in red:
            unflagged.append(p['id'])
        if flags:
            ws.merge_cells('G2:H2')
            tag = ws['G2']
            tag.value = '⚠ ' + '・'.join(flags)
            tag.font = f(13, True, 'FFFFFF')
            tag.alignment = Alignment(horizontal='center', vertical='center')
            for col in (7, 8):
                ws.cell(row=2, column=col).fill = fill('C8372D')

        ws.row_dimensions[3].height = 26
        ws.merge_cells('B3:H3')
        ws['B3'] = p['summary'] + (f"　・　{p['together']}" if p.get('together') else '')
        ws['B3'].font = f(12, False, SOFT)
        ws['B3'].alignment = Alignment(vertical='center')
        ws.row_dimensions[4].height = 18

        # 客戶交代：原文一行一行照抄（她 10/23：「直接呈現所有原文就好」），你的提醒字小一號接在後面
        row = 5
        raw_lines = [t for t in p.get('raw', '').split('\n') if t.strip()]
        reminders = p.get('reminders', [])
        BAR = Side(style='thick', color='3B5BDB')
        lines = [('客戶交代（原文）', 'head')] + [(t, 'raw') for t in raw_lines] + [(t, 'remind') for t in reminders]
        for text, kind in lines:
            size = {'head': 10, 'raw': 13, 'remind': 11}[kind]
            base = {'head': 22, 'raw': 26, 'remind': 24}[kind]
            ws.row_dimensions[row].height = base + (lines_needed(text, size, BOX_WIDTH) - 1) * (size + 6)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            c = ws.cell(row=row, column=2, value=text)
            c.font = {'head': f(10, True, '3B5BDB'), 'raw': f(13, True, '1E2F6B'), 'remind': f(11, False, '1E2F6B')}[kind]
            c.alignment = Alignment(vertical='center', indent=1, wrap_text=True)
            for col in COLS:
                ws.cell(row=row, column=col).fill = fill('DFE6FB')
            ws.cell(row=row, column=2).border = Border(left=BAR)
            row += 1
        ws.row_dimensions[row].height = 22
        row += 1

        ws.row_dimensions[row].height = 24
        for i, c in enumerate(COLS):
            cell = ws.cell(row=row, column=c, value='一二三四五六日'[i])
            cell.font = f(10, True, SOFT)
            cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
            cell.border = Border(bottom=Side(style='thin', color='CFC9BE'))
        ws.row_dimensions[row + 1].height = 8
        row += 2

        for w in range(weeks):
            num_r, lab_r = row, row + 1
            ws.row_dimensions[num_r].height = 24
            ws.row_dimensions[lab_r].height = 34
            for i, c in enumerate(COLS):
                day = first_monday + datetime.timedelta(days=w * 7 + i)
                if day.month != month:
                    continue
                kind = cells[day.day]
                bg, fg, _, _ = LOOK[kind]
                n = ws.cell(row=num_r, column=c, value=day.day)
                lab = ws.cell(row=lab_r, column=c, value=word_of(p, kind, day.day, year, month) or None)
                n.font = f(12, kind != 'notMentioned', fg)
                n.alignment = Alignment(horizontal='left', vertical='bottom', indent=1)
                lab.font = f(11, True, fg)
                lab.alignment = Alignment(horizontal='left', vertical='center', indent=1)
                n.fill = lab.fill = fill(bg)
                n.border = Border(left=GAP, right=GAP, top=GAP)
                lab.border = Border(left=GAP, right=GAP, bottom=GAP)
            row += 2

        # 圖例：只列這一頁用得到的
        row += 1
        ws.row_dimensions[row - 1].height = 18
        seen = []
        for d in range(1, last + 1):
            bg, fg, _, legend = LOOK[cells[d]]
            if (bg, fg, legend) not in seen:
                seen.append((bg, fg, legend))
        seen.sort(key=lambda k: LEGEND_ORDER.index(k[2]))
        ws.row_dimensions[row].height = 24
        col = 2
        for bg, fg, legend in seen:
            span = 1 if len(legend) <= 5 else 2
            if col + span - 1 > 8:
                row += 1
                ws.row_dimensions[row].height = 24
                col = 2
            if span > 1:
                ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
            c = ws.cell(row=row, column=col, value=legend)
            c.font = f(10, True, fg)
            c.alignment = Alignment(horizontal='center', vertical='center')
            for k in range(span):
                ws.cell(row=row, column=col + k).border = Border(left=GAP, right=GAP, top=GAP, bottom=GAP)
                ws.cell(row=row, column=col + k).fill = fill(bg)
            col += span

        if p.get('pending'):
            row += 2
            ws.row_dimensions[row].height = 24
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            ws.cell(row=row, column=2, value=f"還沒確定　{p['pending']}").font = f(10, True, '8A6500')

        # 紙色底：沒有上色的格子全部塗上（格子之間的縫也是這個顏色）
        bottom = row + 2
        for r in range(1, bottom + 1):
            for c in range(1, 10):
                cell = ws.cell(row=r, column=c)
                if cell.fill is None or cell.fill.fill_type != 'solid':
                    cell.fill = fill(PAPER)
        ws.print_area = f'A1:I{bottom}'
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True

    out = os.path.join(folder, f"客戶時間-{plan['month']}.xlsx")
    wb.active = 0
    wb.save(out)
    flagged = sum(1 for p in plan['people'] if (red or {}).get(p['name']))
    with_chart = sum(1 for p in plan['people'] if (charts or {}).get(p['name']) and not charts[p['name']][1])
    odd_chart = sum(1 for p in plan['people'] if (charts or {}).get(p['name']) and charts[p['name']][1])
    print(f"寫好了：{len(plan['people'])} 頁 → {os.path.basename(out)}・病歷號 {with_chart} 位"
          + (f"（另有 {odd_chart} 位沒有或不只一個，頁面上用橘字寫）" if odd_chart else '')
          + f"・標了體內金屬 {flagged} 位"
          + (f"・對不到 app 的 {len(unflagged)} 位沒查：{'、'.join('#' + i for i in unflagged)}" if unflagged else ''))


if __name__ == '__main__':
    build(sys.argv[1])
