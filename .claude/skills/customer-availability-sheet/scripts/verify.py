# 核對：把試算表每一頁月曆的格子讀回來，跟「第二遍讀原文」寫的 核對.json 逐格比。
#
#   python verify.py <月份資料夾>
#
# 核對.json 是**另外讀一次原文、直接列日子**寫的（不看讀法.json、只寫數字），所以這裡比的是兩次獨立的閱讀，
# 不是同一份資料比自己。格式：{ "<讀法.json 的 id>": { "no": [..], "partial": [..], "yes": [..], "unsure": [..],
#   "rest": "free" | "notMentioned" } }。上午不行／下午不行／幾點後才行全部寫進 "partial"（顏色一樣）；
# "free"（沒寫到的日子都可以）跟「可以」同一個綠色（她 10/23），所以比的時候當成 yes。
# 另外比體內金屬：客戶脈絡.json 對到剛好一位、身上有「金屬」的，名字旁邊要有紅色的標；沒有的不可以有。
# 只印頁數、id、日期與種類，不印名字。
import sys, json, os, datetime, calendar
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

COLOR = {'F4B9B3': 'no', 'F6D08F': 'partial', 'B9E2C6': 'yes', 'F9E07E': 'unsure', 'E4E0D8': 'notMentioned'}

folder = sys.argv[1]
plan = json.load(open(os.path.join(folder, '過程', '讀法.json'), encoding='utf-8'))
check = json.load(open(os.path.join(folder, '過程', '核對.json'), encoding='utf-8'))
year, month = (int(x) for x in plan['month'].split('-'))
last = calendar.monthrange(year, month)[1]
wb = load_workbook(os.path.join(folder, f"客戶時間-{plan['month']}.xlsx"))

ids = [p['id'] for p in plan['people']]
problems = []
if set(ids) != set(check):
    problems.append(f'讀法.json 與 核對.json 的人不一樣：只在讀法 {sorted(set(ids) - set(check))}，只在核對 {sorted(set(check) - set(ids))}')
if len(wb.worksheets) != len(ids):
    problems.append(f'試算表 {len(wb.worksheets)} 頁，讀法.json {len(ids)} 位')

first = datetime.date(year, month, 1)
first_monday = first - datetime.timedelta(days=first.weekday())
total = bad = 0
for ws, p in zip(wb.worksheets, plan['people']):
    pid = p['id']
    if ws.title != p['name'][:31]:
        problems.append(f'#{pid} 頁名跟讀法.json 不一樣')
    head = next((r for r in range(1, 40) if [ws.cell(row=r, column=c).value for c in range(2, 9)] == list('一二三四五六日')), None)
    if head is None:
        problems.append(f'#{pid} 找不到星期那一列')
        continue
    top = head + 2
    got = {}
    w = 0
    while True:
        week_start = first_monday + datetime.timedelta(days=w * 7)
        if week_start > datetime.date(year, month, last):
            break
        for i in range(7):
            day = week_start + datetime.timedelta(days=i)
            n = ws.cell(row=top + 2 * w, column=2 + i)
            lab = ws.cell(row=top + 1 + 2 * w, column=2 + i)
            if day.month != month:
                if n.value is not None or lab.value is not None:
                    problems.append(f'#{pid} {day} 不是這個月卻有字')
                continue
            if n.value != day.day:
                problems.append(f'#{pid} 星期對不上：{day} 那一格寫 {n.value}')
            bg = (n.fill.start_color.rgb or '')[-6:]
            if (lab.fill.start_color.rgb or '')[-6:] != bg:
                problems.append(f'#{pid} {month}/{day.day} 上下兩格顏色不一樣')
            got[day.day] = COLOR.get(bg, f'?{bg}')
            if got[day.day] == 'no' and lab.value != '不行':
                problems.append(f'#{pid} {month}/{day.day} 是紅的，格子裡卻寫「{lab.value}」')
            if got[day.day] == 'notMentioned' and lab.value:
                problems.append(f'#{pid} {month}/{day.day} 沒有標種類，格子裡卻寫「{lab.value}」')
        w += 1
    exp = check.get(pid)
    if exp is None:
        continue
    rest = exp.get('rest', 'free')
    want = {d: ('yes' if rest == 'free' else rest) for d in range(1, last + 1)}
    for kind in ('yes', 'partial', 'unsure', 'no'):
        for d in exp.get(kind, []):
            want[d] = kind
    for d in range(1, last + 1):
        total += 1
        if got.get(d) != want[d]:
            bad += 1
            problems.append(f'#{pid} {month}/{d}：試算表「{got.get(d)}」，核對「{want[d]}」')

ctx_path = os.path.join(folder, '過程', '客戶脈絡.json')
if os.path.exists(ctx_path):
    metal = {}
    for c in json.load(open(ctx_path, encoding='utf-8'))['people']:
        if len(c['customers']) == 1:
            metal[c['name']] = any('金屬' in str(x) for x in (c['customers'][0].get('flags') or []))
    tagged = 0
    for ws, p in zip(wb.worksheets, plan['people']):
        has = '金屬' in str(ws['G2'].value or '')
        want = metal.get(p['name'])
        if want is None:
            problems.append(f"#{p['id']} 對不到 app 剛好一位，體內金屬沒有核對")
        elif has != want:
            problems.append(f"#{p['id']} 體內金屬：app {'有' if want else '沒有'}，試算表{'有' if has else '沒有'}標")
        tagged += has
    print(f'體內金屬：標了 {tagged} 位，app 上有 {sum(metal.values())} 位')

    # 病歷號：頁面上寫的（可能留著前面的 0）跟 app 認人用的 chartNos 比，前面的 0 不算
    def norm(x):
        return str(x).strip().lstrip('0') or '0'
    nos = {c['name']: [norm(n) for n in (c['customers'][0].get('chartNos') or [])]
           for c in json.load(open(ctx_path, encoding='utf-8'))['people'] if len(c['customers']) == 1}
    shown = 0
    for ws, p in zip(wb.worksheets, plan['people']):
        want = nos.get(p['name'])
        if want is None:
            continue
        page = str(ws['E2'].value or '')
        if not want:
            if page != 'app 上沒有':
                problems.append(f"#{p['id']} 病歷號：app 上沒有，頁面上寫的不是「app 上沒有」")
        elif sorted(norm(x) for x in page.split('／')) != sorted(want):
            problems.append(f"#{p['id']} 病歷號：頁面上的跟 app 上的不一樣")
        else:
            shown += 1
    print(f"病歷號：{shown} 位跟 app 一樣、{sum(1 for v in nos.values() if not v)} 位 app 上沒有")
else:
    problems.append('沒有 過程/客戶脈絡.json —— 體內金屬沒有核對（先跑 history.mjs）')

for line in problems:
    print(line)
print(f'{len(wb.worksheets)} 頁，比了 {total} 格，對不上 {bad} 格' + ('' if not problems else f'，另有 {len(problems) - bad} 件版面問題' if len(problems) > bad else ''))
sys.exit(1 if problems else 0)
