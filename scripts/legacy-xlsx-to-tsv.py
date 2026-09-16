#!/usr/bin/env python3
"""舊試算表 xlsx → 去識別化的 TSV 樣本。

為什麼要有這支：`domain/legacyImport.js` 是照著 `docs/legacy/README.md` 的
「舊表應該長這樣」寫出來的，從來沒有看過真的檔案。要驗收解析器就得有真的
工作表當樣本 —— 但真的工作表裡有客戶姓名與健康資訊，那些東西進了 git
歷史就拿不掉（SPEC 第 10 節）。

所以這支腳本是那道閘門：**原檔留在 repo 外面，只有去識別化的結果進得來。**

判準是「留下結構、拿掉人」：

  留下   療程名稱、應有／實際次數、日期表頭、勾選格、營養點滴品項明細、
         器材簡寫、TODO 區塊
  拿掉   客戶姓名（→ 客戶A…）、病歷號（→ 9001…）、同事與家屬姓名、
         所有手寫註記（健康狀況與生活細節都在那裡）

寧可多抹一格也不要漏一格：抹掉的是註記，而註記本來就不影響匯入器讀什麼。
輸出之後請自己看過一遍再 commit —— 這支腳本擋的是慣例，不是保證。

用法（原檔請放在 repo 外面）：
    python3 scripts/legacy-xlsx-to-tsv.py <原檔.xlsx> <輸出資料夾>
"""

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('需要 openpyxl：pip install openpyxl')

SCRUBBED = '（手寫註記，已移除）'
THERAPIST = '治A'

# 這些是療程、器材與系統的寫法，不是人。
SAFE_WORDS = re.compile(
    r'高能|sis|SIS|INDIBA|indiba|超磁|雷射|ILIB|EECP|Inbody|復能|靜脈|健檢|'
    r'營養|物理|體適能|復健|心臟|門診|點滴|諮詢|二返|Abovee|Examine|耀聖|打電話|'
    r'TODO|FINISH'
)
CJK = re.compile(r'[\u4e00-\u9fff]')
# A、B 欄裡唯一不是手寫的東西：Apps Script 產生的區塊標題與任務名稱
SYSTEM_WORDS = {'TODO', 'FINISH', 'Abovee', 'Examine', '耀聖', '打電話'}
# 購買名稱長這樣：`0522 顧客會-8`、`0806 H2U導客 -`、`0723顧客會12`
CHANNEL = re.compile(r'^\s*(\d{4})\s*[-\s]*((?:H2U)?\s*(?:顧客會|導客))\s*(.*)$')
# 購買名稱尾巴允許留下來的只有數量與金額，例 `-8+8`、`12`、`-0.75`
QTY_TAIL = re.compile(r'^[-\s0-9+.]*$')
DATE_ONLY = re.compile(r'^\s*\d{1,2}\s*[/月]\s*\d{1,2}\s*日?\s*$')
NUMBER = re.compile(r'^-?\d+(\.\d+)?$')

# 療程列與勾選格的範圍
FIRST_ITEM_ROW, LAST_ITEM_ROW = 2, 12
COL_A, COL_B, COL_C, COL_E = 1, 2, 3, 5
FIRST_DATE_COL = 6


def cell_text(cell):
    """讀成她貼上來的樣子：日期照顯示格式、勾選框照 TRUE/FALSE、整數不帶 .0。"""
    v = cell.value
    if v is None:
        return ''
    if v is True or v is False:
        return 'TRUE' if v else 'FALSE'
    if hasattr(v, 'strftime'):
        # 表頭的日期在畫面上顯示成「6月15日」—— 沒有年份，貼過來也是這樣
        return f'{v.month}月{v.day}日'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def scrub_name_cell(raw, alias, serial):
    """A2 客戶名稱。姓名與病歷號換掉，形狀留著。

    真實的 A2 有三種寫法：`名字`、`名字1234`、`名字\\n(高能/sis)1234`。
    第三種是器材偏好被塞進姓名格 —— 那個括號要留，因為「解析器會不會把它
    當成名字的一部分」正是要測的事。
    """
    keep = ''
    for paren in re.findall(r'[(（]([^)）]*)[)）]', raw):
        if SAFE_WORDS.search(paren):
            keep = f'({paren.strip()})'
        # 括號裡是人（家屬、同事）的一律不留，連括號都不留
    tail = str(serial) if re.search(r'\d{3,}', raw) else ''
    newline = '\n' if '\n' in raw else ''
    return f'{alias}{newline}{keep}{tail}'


def scrub_source_cell(raw):
    """B2 購買名稱。前面的「日期＋通路＋數量」是活動代號，留；後面接的字是註記，抹。"""
    m = CHANNEL.match(raw.replace('\n', ' '))
    if not m:
        return SCRUBBED
    date, channel, tail = m.groups()
    tail = tail.strip()
    if tail and not QTY_TAIL.match(tail):
        # 例：`0604 顧客會-手有金屬，只能INDIBA(某某她最愛)` —— 醫療禁忌加同事名字
        return f'{date} {channel}-{SCRUBBED}'
    return f'{date} {channel}{(" " + tail) if tail else ""}'.rstrip()


def scrub_followup(raw):
    """二返註記：`7/17 二返(夏)` 括號裡是同事，`二返(8/5）` 括號裡是日期。"""
    def one(m):
        inner = m.group(1).strip()
        return f'({inner})' if not inner or DATE_ONLY.match(inner) else f'({THERAPIST})'
    return re.sub(r'[(（]([^)）]*)[)）]', one, raw)


def scrub(row, col, raw, alias, serial, iv_names):
    if not raw.strip():
        return raw
    # 勾選框、次數、日期本身不可能是人
    if raw in ('TRUE', 'FALSE') or NUMBER.match(raw) or DATE_ONLY.match(raw):
        return raw
    if row == 1:
        return raw
    if row == 2 and col == COL_A:
        return scrub_name_cell(raw, alias, serial)
    if row == 2 and col == COL_B:
        return scrub_source_cell(raw)
    # 二返註記寫在日期欄的第 13 列。A、B 欄裡出現的「二返」是手寫的，
    # 例如「二返X光+自然美」—— 那是檢查項目，走下面的註記規則抹掉。
    if '二返' in raw and col >= FIRST_DATE_COL:
        return scrub_followup(raw)
    # 療程列的 B～E 欄是結構：B 是營養點滴的品項明細，C～E 是療程名稱與次數
    if FIRST_ITEM_ROW <= row <= LAST_ITEM_ROW and COL_B <= col <= COL_E:
        return raw
    # 日期欄。勾選格已經在上面回去了，剩下的中文多半是手寫的，
    # 例外是「當天用了哪個品項」的簡寫 —— 它對得上 B 欄的品項名，留著。
    if col >= FIRST_DATE_COL:
        if SAFE_WORDS.search(raw) or not CJK.search(raw):
            return raw
        if any(n.startswith(raw) or raw.startswith(n) for n in iv_names):
            return raw
        return SCRUBBED
    # 剩下的是 A、B 欄的空白處 —— 手寫註記都在這裡，所以這一區反過來：
    # 只有「認得出是系統產生的」才留，其餘一律抹掉。
    if raw.strip() in SYSTEM_WORDS:
        return raw
    # TODO 區塊的「6/15復能(1小時)」。日期後面要接得出療程名才留，否則
    # 「8/4核磁共振」這種手寫的檢查紀錄會混在裡面被當成任務列放行。
    if re.match(r'^\s*\d{1,2}/\d{1,2}', raw) and SAFE_WORDS.search(raw):
        return raw
    return SCRUBBED


def iv_product_names(ws):
    """B11、B12 的品項明細裡出現過的名字。第 14 列的簡寫要對照它們才留得下來。"""
    names = []
    for row in (11, 12):
        for piece in re.split(r'[+＋、,，]', str(ws.cell(row, COL_B).value or '')):
            name = re.split(r'[x×*]', piece.strip())[0].strip()
            if name:
                names.append(name)
    return names


def convert(src, out_dir):
    wb = openpyxl.load_workbook(src, data_only=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []
    serial = 9000
    nth = 0

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        if ws.sheet_state != 'visible':
            continue
        if '模板' in sheet:
            alias = name = sheet
        else:
            serial += 1
            alias = name = f'客戶{chr(ord("A") + nth)}'
            nth += 1

        iv_names = iv_product_names(ws)
        rows = []
        for r in range(1, min(ws.max_row, 60) + 1):
            cells = [scrub(r, c, cell_text(ws.cell(r, c)), alias, serial, iv_names)
                     for c in range(1, min(ws.max_column, 90) + 1)]
            while cells and cells[-1] == '':
                cells.pop()
            rows.append(cells)
        while rows and not any(c.strip() for c in rows[-1]):
            rows.pop()

        # 儲存格裡的換行會把 TSV 拆成兩列，改成 \n 的字面寫法留著形狀
        text = '\n'.join('\t'.join(c.replace('\n', '\\n').replace('\t', ' ') for c in row)
                         for row in rows)
        (out_dir / f'{name}.tsv').write_text(text + '\n', encoding='utf-8')
        made.append(name)
        print(f'{sheet!r:24} → {name}.tsv  {len(rows)} 列')

    return made


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
